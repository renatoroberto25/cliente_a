#!/usr/bin/env python3
"""Gera um baseline Excel independente para cada audit COMPLIANCE executado."""

from __future__ import annotations

import hashlib
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = ROOT / "01_Hardening_2026_RHEL_LIKE_PADRAO.xlsx"

HEADERS = [
    "Descrição",
    "Perfil",
    "ID",
    "Item",
    "Justificativa",
    "Auditoria - Descrição",
    "Auditoria - Script",
    "Correção - Descrição",
    "Correção - Script",
    "Criticidade",
    "Impacto",
]

FIELD_RE = re.compile(r'^\s*([A-Za-z_]+)\s*:\s*"(.*)"\s*$', re.MULTILINE)
ITEM_RE = re.compile(r"<custom_item>(.*?)</custom_item>", re.DOTALL | re.IGNORECASE)
DESCRIPTION_RE = re.compile(r"^\s*(\d+)\s*,\s*(.*?)\s*,\s*([^,]+)\s*$")
DESCRIPTION_WITHOUT_SEVERITY_RE = re.compile(r"^\s*(\d+)\s*,\s*(.+?)\s*$")
WINDOWS_DESCRIPTION_RE = re.compile(r"^\s*(\d+)\s+(.*?)\s+(Alta|M[eé]dia|Baixa)\s*$", re.IGNORECASE)


def audit_items(path: Path) -> list[dict[str, str]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    items = []
    for position, block in enumerate(ITEM_RE.findall(text), 1):
        fields = {key.lower(): value for key, value in FIELD_RE.findall(block)}
        description = fields.get("description", "")
        match = DESCRIPTION_RE.match(description)
        match_without_severity = DESCRIPTION_WITHOUT_SEVERITY_RE.match(description)
        windows_match = WINDOWS_DESCRIPTION_RE.match(description)
        if match:
            item_id, title, severity = match.groups()
        elif match_without_severity:
            item_id, title = match_without_severity.groups()
            severity = "Não informada"
        elif windows_match:
            item_id, title, severity = windows_match.groups()
        else:
            item_id = f"AUTO-{position}"
            title = description or f"Controle {position}"
            severity = "Não informada"
        reference = fields.get("reference", "")
        profile_match = re.search(r",\s*(Full|Light|Manual)\s*$", reference, re.IGNORECASE)
        excluded = {"description", "info", "solution", "reference", "see_also"}
        technical = " ".join(f"{key}: {value}" for key, value in fields.items() if key not in excluded)
        items.append(
            {
                "id": item_id,
                "title": title.strip(),
                "severity": severity.strip(),
                "profile": profile_match.group(1).title() if profile_match else "Full",
                "info": fields.get("info", ""),
                "solution": fields.get("solution", ""),
                "cmd": fields.get("cmd", ""),
                "expect": fields.get("expect", ""),
                "type": fields.get("type", ""),
                "technical": technical,
            }
        )
    return items


def fingerprint(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def classify(name: str) -> str:
    upper = name.upper()
    if "WINDOWS" in upper:
        return "Windows Server 2022"
    if "DEB" in upper:
        return "Debian-like"
    if "RHEL" in upper:
        return "RHEL 7-9"
    return "Utilitário"


def impact_for(title: str, platform: str) -> str:
    value = title.lower()
    if any(word in value for word in ("ssh", "rede", "firewall", "porta", "protocolo")):
        return "Pode afetar conectividade ou administração remota; validar acesso alternativo antes da remediação."
    if any(word in value for word in ("módulo", "modulo", "kernel", "usb", "filesystem")):
        return "Pode afetar recurso de kernel ou periférico em uso; confirmar dependências antes de bloquear."
    if platform.startswith("Windows"):
        return "Pode afetar serviços, políticas ou aplicações do Windows; validar em homologação antes da remediação."
    return "Pode afetar serviços ou aplicações; validar em homologação e registrar exceções antes da remediação."


def copy_cell_style(source, target) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def prepare_control_sheet(template_ws, title: str, items: list[dict[str, str]], platform: str):
    ws = template_ws.parent.copy_worksheet(template_ws)
    ws.title = title
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for column, header in enumerate(HEADERS, 1):
        ws.cell(1, column, header)
    for row_number, item in enumerate(items, 2):
        solution = item["solution"] or "Definir e validar a remediação correspondente ao controle."
        audit_description = (
            f"Executar o controle {item['type'] or 'de auditoria'} conforme o comando utilizado no scan. "
            f"Resultado esperado: {item['expect'] or 'conforme definido no audit'}."
        )
        values = [
            item["title"],
            item["profile"],
            int(item["id"]) if item["id"].isdigit() else item["id"],
            item["title"],
            item["info"],
            audit_description,
            item["technical"],
            solution,
            f"Aplicar: {solution} Validar novamente pelo comando do item e registrar evidência.",
            item["severity"],
            impact_for(item["title"], platform),
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            copy_cell_style(template_ws.cell(2, column), cell)
            cell.alignment = copy(template_ws.cell(2, column).alignment)
    ws.auto_filter.ref = f"A1:K{max(1, len(items) + 1)}"
    ws.freeze_panes = "A2"
    return ws


def style_index(ws) -> None:
    headers = ["Aba", "Plataforma", "Controles", "SHA-256", "Audit representante", "Arquivos equivalentes/origens"]
    for column, value in enumerate(headers, 1):
        cell = ws.cell(1, column, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="FFFFFF"))
    widths = [14, 22, 12, 68, 62, 95]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def main() -> None:
    if not TEMPLATE.is_file():
        raise SystemExit(
            f"Template obrigatório não encontrado: {TEMPLATE}\n"
            "Coloque 01_Hardening_2026_RHEL_LIKE_PADRAO.xlsx um nível acima da pasta do script."
        )

    audits = sorted(HERE.glob("COMPLIANCE_*.audit"))
    if not audits:
        raise SystemExit("Nenhum arquivo .audit encontrado.")

    for audit in audits:
        output_book = load_workbook(TEMPLATE)
        template_ws = output_book.active
        items = audit_items(audit)
        ids = [item["id"] for item in items]
        duplicate_ids = sorted({item_id for item_id in ids if ids.count(item_id) > 1})
        if duplicate_ids:
            raise SystemExit(f"IDs duplicados em {audit.name}: {', '.join(duplicate_ids)}")
        platform = classify(audit.name)
        baseline_ws = prepare_control_sheet(template_ws, "Baseline", items, platform)
        output_book.remove(template_ws)

        metadata = output_book.create_sheet("Origem", 0)
        metadata.append(["Campo", "Valor"])
        metadata.append(["Audit de origem", audit.name])
        metadata.append(["Plataforma", platform])
        metadata.append(["Controles", len(items)])
        metadata.append(["SHA-256", fingerprint(audit)])
        style_index(metadata)
        metadata.delete_cols(3, 4)
        metadata.column_dimensions["A"].width = 24
        metadata.column_dimensions["B"].width = 100
        metadata.auto_filter.ref = "A1:B5"

        output = HERE / f"BASELINE_{audit.stem}.xlsx"
        output_book.save(output)
        print(f"Gerado: {output.name} ({len(items)} controles)")

    print(f"Total: {len(audits)} audits; {len(audits)} baselines independentes")


if __name__ == "__main__":
    main()
