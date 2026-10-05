#!/usr/bin/env python3
"""Gera o audit RHEL6/OL6 a partir da planilha legada oficial do projeto."""

from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "03_Hardening_2026_RHEL_LEGADO.xlsx"
OUTPUT = ROOT / "COMPLIANCE_CLARO_RHEL6_LEGADO_CUSTOM.audit"


def ok(test):
    return "( " + test + " ) && echo PASS || echo FAIL"


def absent_pkg(*pkgs):
    return ok("! rpm -q " + " ".join(pkgs) + " >/dev/null 2>&1")


def service_off(*names):
    checks = []
    for name in names:
        checks += [f"! service {name} status >/dev/null 2>&1", f"! chkconfig --list {name} 2>/dev/null | grep -Eq '[2-5]:on'"]
    return ok(" && ".join(checks))


def sysctl_values(pairs):
    return ok(" && ".join(f"[ \"$(sysctl -n {key} 2>/dev/null)\" = \"{value}\" ]" for key, value in pairs))


def mountpoint(path):
    return ok(f"awk '$2==\"{path}\"{{found=1}} END{{exit !found}}' /proc/mounts")


def mountopts(path, options):
    req = " && ".join(f"echo \"$opts\" | grep -Eq '(^|,){x}(,|$)'" for x in options)
    return f"opts=$(awk '$2==\"{path}\"{{print $4; exit}}' /proc/mounts); ( [ -n \"$opts\" ] && {req} ) && echo PASS || echo FAIL"


def ssh(setting, value_re):
    # sshd -T existe nas versões RHEL6 suportadas; fallback lê a configuração.
    return ok(f"( sshd -T 2>/dev/null || sed -e 's/#.*//' /etc/ssh/sshd_config ) | grep -Eiq '^[[:space:]]*{setting}[[:space:]]+{value_re}([[:space:]]|$)'")


def audit_rule(pattern):
    return ok(f"grep -ERq -- '{pattern}' /etc/audit/audit.rules /etc/audit/rules.d 2>/dev/null")


CMDS = {}

# Kernel: sintaxe compatível com module-init-tools e grep ERE do EL6.
for ident, module in {1:'cramfs',2:'squashfs',3:'udf',4:'hfs',5:'hfsplus',6:'jffs2',7:'freevxfs',8:'overlay',9:'usb_storage',10:'dccp',11:'sctp',12:'rds',13:'tipc'}.items():
    CMDS[ident] = ok(f"( ! modinfo {module} >/dev/null 2>&1 ) || ( modprobe -n -v {module} 2>/dev/null | grep -Eq '^[[:space:]]*install[[:space:]]+/bin/(true|false)([[:space:]]|$)' && ! lsmod | awk '{{print $1}}' | grep -qx '{module}' )")
CMDS.update({
    14: sysctl_values([('fs.suid_dumpable', '0')]),
    15: sysctl_values([('kernel.randomize_va_space', '2')]),
    16: ok("[ \"$(ulimit -c)\" = 0 ] && [ \"$(sysctl -n fs.suid_dumpable 2>/dev/null)\" = 0 ]"),
    17: ok("grep -Eq '(^|[[:space:]])nx([[:space:]]|$)' /proc/cpuinfo || dmesg 2>/dev/null | grep -Eiq 'NX.*(active|enabled|protection)'"),
    18: ok("[ ! -e /proc/sys/kernel/perf_event_paranoid ] || [ \"$(cat /proc/sys/kernel/perf_event_paranoid)\" -ge 2 ]"),
    19: ok("[ ! -e /proc/sys/kernel/dmesg_restrict ] || [ \"$(cat /proc/sys/kernel/dmesg_restrict)\" = 1 ]"),
    20: ok("[ ! -e /proc/sys/fs/protected_symlinks ] || [ \"$(cat /proc/sys/fs/protected_symlinks)\" = 1 ]"),
    21: ok("[ ! -e /proc/sys/fs/protected_hardlinks ] || [ \"$(cat /proc/sys/fs/protected_hardlinks)\" = 1 ]"),
})

for ident, path in {22:'/tmp',23:'/dev/shm',24:'/var',25:'/var/tmp',26:'/var/log',27:'/var/log/audit',28:'/home'}.items():
    CMDS[ident] = mountpoint(path)
for ident, spec in {29:('/tmp',['nodev','nosuid','noexec']),30:('/dev/shm',['nodev','nosuid','noexec']),31:('/var/tmp',['nodev','nosuid','noexec']),32:('/var/log',['nodev','nosuid']),33:('/var/log/audit',['nodev','nosuid']),34:('/home',['nodev'])}.items():
    CMDS[ident] = mountopts(*spec)
CMDS.update({
    35: ok("! find / -xdev -type d -perm -0002 ! -perm -1000 -print -quit 2>/dev/null | grep -q ."),
    36: ok("[ -d /tmp ] && [ \"$(stat -c %a /tmp)\" = 1777 ]"),
    37: ok("[ -d /var/tmp ] && [ \"$(stat -c %a /var/tmp)\" = 1777 ]"),
    38: ok("[ -d /dev/shm ] && [ \"$(stat -c %a /dev/shm)\" = 1777 ]"),
    39: service_off('autofs'),
    40: ok("grep -Eq '^[[:space:]]*password[[:space:]]+--(md5|encrypted)[[:space:]]+' /boot/grub/grub.conf"),
    41: ok("stat -c '%a %u %g' /boot/grub/grub.conf 2>/dev/null | grep -Eq '^(400|600) 0 0$'"),
    42: ok("grep -Eq '^[[:space:]]*SINGLE=/sbin/sulogin([[:space:]]|$)' /etc/sysconfig/init"),
    43: absent_pkg('prelink'), 44: absent_pkg('xinetd'),
    45: ok("service ntpd status >/dev/null 2>&1"),
    46: ok("grep -Eq '^[[:space:]]*(server|pool)[[:space:]]+[^[:space:]#]+' /etc/ntp.conf"),
    47: ok("chkconfig --list ntpd 2>/dev/null | grep -Eq '[2-5]:on'"),
    48: ok("service ntpd status >/dev/null 2>&1 && ! pgrep -x chronyd >/dev/null 2>&1"),
    49: absent_pkg('xorg-x11-server-Xorg'), 50: service_off('avahi-daemon'), 51: service_off('cups'),
    52: absent_pkg('dhcp'), 53: absent_pkg('openldap-servers'), 54: absent_pkg('bind'),
    55: absent_pkg('vsftpd'), 56: absent_pkg('httpd'), 57: absent_pkg('dovecot','cyrus-imapd'),
    58: absent_pkg('samba'), 59: absent_pkg('squid'), 60: service_off('snmpd'), 61: absent_pkg('ypserv'),
    62: absent_pkg('telnet-server'),
    63: ok("! netstat -lntp 2>/dev/null | awk '$4 ~ /:25$/ && $4 !~ /^(127\\.0\\.0\\.1|::1):25$/ {bad=1} END{exit !bad}'"),
    64: service_off('nfs','nfslock'), 65: service_off('rpcbind','portmap'),
    66: ok("! pgrep -x rsync >/dev/null 2>&1 && ! grep -ERq '^[[:space:]]*disable[[:space:]]*=[[:space:]]*no' /etc/xinetd.d/rsync 2>/dev/null"),
    67: absent_pkg('ypbind'), 68: absent_pkg('rsh'), 69: absent_pkg('talk'), 70: absent_pkg('telnet'),
    71: absent_pkg('openldap-clients'),
    72: ok("! chkconfig --list 2>/dev/null | grep -E '^(avahi-daemon|cups|vsftpd|xinetd|tftp|ypserv|ypbind)[[:space:]].*[2-5]:on'"),
    73: absent_pkg('tftp-server'),
    74: ok("awk -F= '/^[[:space:]]*gpgcheck[[:space:]]*=/{gsub(/[[:space:]]/,\"\",$2); if($2!=1) bad=1; seen=1} END{exit bad || !seen}' /etc/yum.conf /etc/yum.repos.d/*.repo 2>/dev/null"),
    75: ok("yum -q check-update --security >/dev/null 2>&1; rc=$?; [ $rc -eq 0 ]"),
    76: ok("[ -f /etc/cron.allow ] && [ ! -s /etc/cron.deny ]"),
    77: ok("[ -f /etc/at.allow ] && [ ! -s /etc/at.deny ]"),
})

CMDS.update({
    78: sysctl_values([('net.ipv4.ip_forward','0')]),
    79: sysctl_values([('net.ipv4.conf.all.send_redirects','0'),('net.ipv4.conf.default.send_redirects','0')]),
    80: sysctl_values([('net.ipv4.conf.all.accept_source_route','0'),('net.ipv4.conf.default.accept_source_route','0')]),
    81: sysctl_values([('net.ipv4.conf.all.accept_redirects','0'),('net.ipv4.conf.default.accept_redirects','0'),('net.ipv4.conf.all.secure_redirects','0'),('net.ipv4.conf.default.secure_redirects','0')]),
    82: sysctl_values([('net.ipv4.conf.all.rp_filter','1'),('net.ipv4.conf.default.rp_filter','1')]),
    83: sysctl_values([('net.ipv4.icmp_echo_ignore_broadcasts','1')]),
    84: sysctl_values([('net.ipv4.icmp_ignore_bogus_error_responses','1')]),
    85: sysctl_values([('net.ipv4.tcp_syncookies','1')]),
    86: ok("[ \"$(cat /proc/sys/net/ipv6/conf/all/disable_ipv6 2>/dev/null)\" = 1 ] || ( [ \"$(sysctl -n net.ipv6.conf.all.accept_redirects 2>/dev/null)\" = 0 ] && [ \"$(sysctl -n net.ipv6.conf.default.accept_redirects 2>/dev/null)\" = 0 ] )"),
    87: ok("! find /sys/class/net -maxdepth 2 -type d -name wireless -print -quit 2>/dev/null | grep -q . && ! hciconfig 2>/dev/null | grep -q '^hci'"),
    93: ok("stat -c '%a %u %g' /etc/ssh/sshd_config 2>/dev/null | grep -Eq '^(600|640) 0 0$'"),
    94: ssh('protocol','2'),
    95: ok("grep -Eiq '^[[:space:]]*(AllowUsers|AllowGroups|DenyUsers|DenyGroups)[[:space:]]+' /etc/ssh/sshd_config"),
    96: ssh('syslogfacility','AUTHPRIV'), 97: ssh('loglevel','(INFO|VERBOSE)'),
    98: ssh('x11forwarding','no'), 99: ssh('maxauthtries','[1-4]'), 100: ssh('ignorerhosts','yes'),
    101: ssh('hostbasedauthentication','no'), 102: ssh('permitrootlogin','(no|without-password|forced-commands-only)'),
    103: ssh('permitemptypasswords','no'), 104: ssh('permituserenvironment','no'), 105: ssh('usepam','yes'),
    106: ok("v=$(sshd -T 2>/dev/null); echo \"$v\" | grep -Eq '^clientaliveinterval [1-9][0-9]*$' && echo \"$v\" | grep -Eq '^clientalivecountmax [0-3]$'"),
    107: ssh('logingracetime','([1-5]?[0-9]|60)'),
    108: ssh('maxstartups','(10|[1-9]):[0-9]+:[0-9]+'), 109: ssh('allowtcpforwarding','no'),
    110: ok("grep -Eiq '^[[:space:]]*Banner[[:space:]]+/[^[:space:]#]+' /etc/ssh/sshd_config"),
    111: ok("! grep -Eiq '^[[:space:]]*Ciphers[[:space:]].*(arcfour|3des|cbc)' /etc/ssh/sshd_config"),
    112: ok("! grep -Eiq '^[[:space:]]*MACs[[:space:]].*(hmac-md5|hmac-sha1-96|hmac-md5-96)' /etc/ssh/sshd_config"),
    113: ok("! sshd -T 2>&1 | grep -qi 'Bad configuration option: KexAlgorithms' || ! grep -Eiq '^[[:space:]]*KexAlgorithms[[:space:]].*diffie-hellman-group1-sha1' /etc/ssh/sshd_config"),
})

CMDS.update({
    114: ok("grep -Eq 'pam_(cracklib|pwquality)\\.so' /etc/pam.d/system-auth /etc/pam.d/password-auth 2>/dev/null"),
    115: ok("grep -E 'pam_cracklib\\.so' /etc/pam.d/system-auth /etc/pam.d/password-auth 2>/dev/null | grep -Eq 'minlen=1[4-9]|minlen=[2-9][0-9]'"),
    116: ok("grep -Eq 'pam_tally2\\.so.*deny=[1-5]' /etc/pam.d/system-auth /etc/pam.d/password-auth 2>/dev/null"),
    117: ok("grep -Eq 'pam_unix\\.so.*remember=([5-9]|[1-9][0-9]+)' /etc/pam.d/system-auth /etc/pam.d/password-auth 2>/dev/null"),
    118: ok("grep -Eq '^[[:space:]]*ENCRYPT_METHOD[[:space:]]+SHA512' /etc/login.defs && grep -Eq 'pam_unix\\.so.*sha512' /etc/pam.d/system-auth 2>/dev/null"),
    119: ok("awk '$1==\"PASS_MAX_DAYS\"{exit !($2>0 && $2<=90)} END{if(NR==0)exit 1}' /etc/login.defs"),
    120: ok("awk '$1==\"PASS_MIN_DAYS\"{exit !($2>=1)}' /etc/login.defs"),
    121: ok("awk '$1==\"PASS_WARN_AGE\"{exit !($2>=7)}' /etc/login.defs"),
    122: ok("useradd -D 2>/dev/null | grep -Eq '^INACTIVE=-?[0-9]+' && ! useradd -D 2>/dev/null | grep -q '^INACTIVE=-1$'"),
    123: ok("! awk -F: -v now=$(date +%s) '$3!="" && ($3*86400)>now{print $1}' /etc/shadow | grep -q ."),
    124: ok("[ \"$(awk -F: '$3==0{n++} END{print n+0}' /etc/passwd)\" = 1 ] && awk -F: '$3==0{exit !($1==\"root\")}' /etc/passwd"),
    125: ok("[ \"$(awk -F: '$3==0{n++} END{print n+0}' /etc/group)\" = 1 ] && grep -q '^root:x:0:' /etc/group"),
    126: ok("! awk -F: '$3<500 && $1!~/^(root|sync|shutdown|halt)$/ && $7!~/(nologin|false)$/ {print}' /etc/passwd | grep -q ."),
    127: ok("[ -s /etc/shells ] && ! grep -Ev '^$|^/(bin|sbin|usr/bin)/(ba|c|k|z)?sh$|^/bin/(false|nologin)$|^/sbin/nologin$' /etc/shells | grep -q ."),
    128: ok("p=$(su - root -c 'printf %s \"$PATH\"' 2>/dev/null); [ -n \"$p\" ] && ! echo \"$p\" | grep -Eq '(^|:)(\\.|[^:]*[[:space:]][^:]*|)(:|$)|::$'"),
    129: ok("grep -ERq '^[[:space:]]*umask[[:space:]]+0?27([[:space:]]|$)' /etc/profile /etc/bashrc /etc/profile.d 2>/dev/null"),
    130: ok("grep -ERq '^[[:space:]]*(readonly[[:space:]]+)?TMOUT=([1-8]?[0-9]?[0-9]|900)([[:space:]]|$|;)' /etc/profile /etc/bashrc /etc/profile.d 2>/dev/null"),
    131: ok("visudo -c >/dev/null 2>&1 && [ -f /etc/sudoers ]"),
    132: ok("grep -ERq '^[[:space:]]*Defaults([^#]*,)?requiretty([,[:space:]]|$)' /etc/sudoers /etc/sudoers.d 2>/dev/null"),
    133: ok("grep -ERq '^[[:space:]]*Defaults.*(logfile|syslog)=' /etc/sudoers /etc/sudoers.d 2>/dev/null"),
    134: ok("! grep -ER '^[^#].*(NOPASSWD|!authenticate)' /etc/sudoers /etc/sudoers.d 2>/dev/null | grep -q ."),
    135: ok("grep -ERq '^[[:space:]]*Defaults.*timestamp_timeout=([0-9]|1[0-5])([,[:space:]]|$)' /etc/sudoers /etc/sudoers.d 2>/dev/null"),
    136: ok("grep -Eq '^[[:space:]]*auth[[:space:]]+required[[:space:]]+pam_wheel\\.so.*use_uid' /etc/pam.d/su"),
})

CMDS.update({
    137: ok("rpm -q rsyslog >/dev/null 2>&1 && service rsyslog status >/dev/null 2>&1"),
    138: ok("grep -Eq '^[[:space:]]*[^#].*(/var/log/|@@?[^[:space:]]+)' /etc/rsyslog.conf /etc/rsyslog.d/*.conf 2>/dev/null"),
    139: ok("rpm -q logrotate >/dev/null 2>&1 && [ -s /etc/logrotate.conf ]"),
    140: ok("rpm -q audit >/dev/null 2>&1"), 141: ok("service auditd status >/dev/null 2>&1"),
    142: ok("chkconfig --list auditd 2>/dev/null | grep -Eq '[2-5]:on'"),
    143: ok("grep -Eq '(^|[[:space:]])audit=1([[:space:]]|$)' /proc/cmdline && grep -Eq '^[[:space:]]*kernel.*audit=1' /boot/grub/grub.conf"),
    144: ok("awk -F= '$1~/^[[:space:]]*max_log_file[[:space:]]*$/{gsub(/[[:space:]]/,\"\",$2); exit !($2>=8)}' /etc/audit/auditd.conf"),
    145: ok("grep -Eiq '^[[:space:]]*max_log_file_action[[:space:]]*=[[:space:]]*(keep_logs|rotate)' /etc/audit/auditd.conf"),
    146: audit_rule('(adjtimex|settimeofday|clock_settime|/etc/localtime)'),
    147: audit_rule('(/etc/passwd|/etc/group|/etc/shadow|/etc/gshadow|/etc/security/opasswd)'),
    148: audit_rule('(/etc/hosts|/etc/sysconfig/network|/etc/sysconfig/network-scripts)'),
    150: audit_rule('(/var/log/faillog|/var/log/lastlog|/var/run/utmp|/var/log/wtmp|/var/log/btmp)'),
    151: audit_rule('(exit=-EACCES|exit=-EPERM)'), 152: audit_rule('perm=x.*(auid>=500|auid!=4294967295)'),
    153: audit_rule('(-S[[:space:]]+mount|,mount)'), 154: audit_rule('(unlink|unlinkat|rename|renameat)'),
    155: audit_rule('(init_module|delete_module|/sbin/insmod|/sbin/rmmod|/sbin/modprobe)'),
    156: audit_rule('(/etc/sudoers|/etc/sudoers.d|/var/log/sudo)'),
    157: ok("grep -Eq '^[[:space:]]*-e[[:space:]]+2[[:space:]]*$' /etc/audit/audit.rules"),
    158: ok("rpm -q aide >/dev/null 2>&1"),
    159: ok("[ -s /var/lib/aide/aide.db.gz ] || [ -s /var/lib/aide/aide.db ]"),
    160: ok("grep -ERq '^[^#].*aide.*--check' /etc/crontab /etc/cron.d /var/spool/cron 2>/dev/null"),
})

def file_mode(path, modes):
    regex = '|'.join(str(x) for x in modes)
    return ok(f"stat -c '%a %u %g' {path} 2>/dev/null | grep -Eq '^({regex}) 0 0$'")

CMDS.update({
    161:file_mode('/etc/passwd',[600,640,644]), 162:file_mode('/etc/passwd-',[600,640,644]),
    163:file_mode('/etc/shadow',[0,400,600,640]), 164:file_mode('/etc/shadow-',[0,400,600,640]),
    165:file_mode('/etc/gshadow',[0,400,600,640]), 166:file_mode('/etc/gshadow-',[0,400,600,640]),
    167:file_mode('/etc/group',[600,640,644]), 168:file_mode('/etc/group-',[600,640,644]),
    169: ok("[ ! -e /etc/security/opasswd ] || stat -c '%a %u %g' /etc/security/opasswd | grep -Eq '^(400|600) 0 0$'"),
    170: ok("! find / -xdev -nouser -print -quit 2>/dev/null | grep -q ."),
    171: ok("! find / -xdev -nogroup -print -quit 2>/dev/null | grep -q ."),
    172: ok("a=$(mktemp /tmp/rhel6_suid.XXXXXX) && b=$(mktemp /tmp/rhel6_suid_rpm.XXXXXX) || exit 1; find / -xdev -type f \\( -perm -4000 -o -perm -2000 \\) -print 2>/dev/null | sort > \"$a\"; rpm -Va 2>/dev/null | awk '$1~/[SM5DLUGTP]/ && $2~/[ug]/ {print $NF}' | sort > \"$b\"; diff -q \"$a\" \"$b\" >/dev/null 2>&1; rc=$?; rm -f \"$a\" \"$b\"; [ $rc -eq 0 ]"),
    173: ok("! awk -F: '$3>=500 && $1!=\"nfsnobody\" && $7!~/(nologin|false)$/ {print $1, $6}' /etc/passwd | while read u h; do [ -d \"$h\" ] || echo \"$u\"; done | grep -q ."),
    174: ok("awk -F: '$3>=500 && $1!=\"nfsnobody\" && $7!~/(nologin|false)$/ {print $1, $6}' /etc/passwd | while read u h; do [ ! -d \"$h\" ] || [ \"$(stat -c %U \"$h\")\" = \"$u\" ] || exit 1; done"),
    175: ok("awk -F: '$3>=500 && $1!=\"nfsnobody\" && $7!~/(nologin|false)$/ {print $6}' /etc/passwd | while read h; do [ ! -d \"$h\" ] || [ $((0$(stat -c %a \"$h\") & 0027)) -eq 0 ] || exit 1; done"),
    176: ok("! awk -F: '$3>=500 && $1!=\"nfsnobody\" && $7!~/(nologin|false)$/ {print $6}' /etc/passwd | while read h; do [ -d \"$h\" ] && find \"$h\" -maxdepth 1 -type f -name '.*' -perm /022 -print; done | grep -q ."),
    177: ok("! awk -F: '$3>=500 && $1!=\"nfsnobody\" && $7!~/(nologin|false)$/ {print $6}' /etc/passwd | while read h; do for f in .forward .netrc .rhosts; do [ -e \"$h/$f\" ] && echo \"$h/$f\"; done; done | grep -q ."),
    178: ok(". /etc/os-release 2>/dev/null || :; rel=$(cat /etc/redhat-release 2>/dev/null); echo \"$rel $VERSION_ID\" | grep -Eiq '(Red Hat Enterprise Linux|CentOS|Oracle Linux|Scientific Linux).*(release[[:space:]]+6|6\\.)|VERSION_ID=.?6'"),
    179: ok("[ -s /etc/rhel6-support-contract.evidence ]"),
    180: ok("rpm -q --qf '%{VENDOR} %{VERSION}-%{RELEASE}\\n' kernel 2>/dev/null | grep -Eiq '(Red Hat|Oracle)' && [ -r /etc/redhat-release ]"),
    181: ok("rpm -q openssl 2>/dev/null | grep -Eq '^openssl-1\\.0\\.1e-'"),
    182: ok("! grep -ERiq '(^|[^A-Za-z0-9])(SSLv2|SSLv3|TLSv1([[:space:]]|$)|RC4|DES-CBC3-SHA)([^A-Za-z0-9]|$)' /etc/httpd /etc/nginx /etc/haproxy /etc/stunnel 2>/dev/null"),
    183: ok("! testparm -s 2>/dev/null | grep -Eiq '^[[:space:]]*(server min protocol|client min protocol)[[:space:]]*=[[:space:]]*(LANMAN1|NT1)'"),
    184: ok("! awk '$1!~/^#/ && ($2~/rw|insecure|no_root_squash/) {bad=1} END{exit !bad}' /etc/exports 2>/dev/null"),
    185: ok("t=$(mktemp /tmp/rhel6_ports.XXXXXX) || exit 1; netstat -lntup 2>/dev/null | awk 'NR>2{print $4}' | sort -u > \"$t\"; [ -s /etc/rhel6-approved-listeners ] && diff -q \"$t\" /etc/rhel6-approved-listeners >/dev/null 2>&1; rc=$?; rm -f \"$t\"; [ $rc -eq 0 ]"),
    186: absent_pkg('gcc','gcc-c++','make'),
    188: ok("yum -q repolist enabled 2>/dev/null | grep -Eiq '(rhel|red hat|oracle|ol6|satellite)'"),
    189: ok("! netstat -lntp 2>/dev/null | grep -E 'java|:1099[[:space:]]|:8009[[:space:]]'"),
    190:file_mode('/etc/gshadow',[0,400,600,640]),
})


def esc(value):
    return str(value).replace('\\', '\\\\').replace('"', '\\"').replace('\n', ' ')


ws = load_workbook(SOURCE, read_only=True, data_only=True).active
rows = list(ws.iter_rows(min_row=2, values_only=True))
ids = [int(row[0]) for row in rows]
missing = [ident for ident in ids if ident not in CMDS]
extra = [ident for ident in CMDS if ident not in ids]
if missing or extra:
    raise SystemExit(f"Mapa inconsistente: sem comando={missing}; comando sem planilha={extra}")

parts = [
    '<check_type:"Unix">',
    '',
    '# Gerado de 03_Hardening_2026_RHEL_LEGADO.xlsx.',
    '# Escopo: Red Hat Enterprise Linux 6.x e Oracle Linux 6.x.',
    '# IDs ausentes na origem permanecem ausentes; não houve renumeração.',
    '',
]
for ident, topic, profile, severity, summary, description, audit_description in rows:
    ident = int(ident)
    parts += [
        '<custom_item>',
        '  type        : CMD_EXEC',
        f'  description : "RHEL6-{ident:03d} | {esc(summary)} | Criticidade: {esc(severity)} | RHEL6 LEGADO"',
        f'  info        : "{esc(audit_description)}. {esc(description)}"',
        f'  solution    : "{esc(description)}"',
        f'  reference   : "03_Hardening_2026_RHEL_LEGADO.xlsx | ID {ident} | {esc(topic)} | Perfil {esc(profile)} | RHEL6 LEGADO"',
        '  see_also    : "Baseline RHEL Legado; validar exceções e suporte contratado"',
        f'  cmd         : "{esc(CMDS[ident])}"',
        '  expect      : "(?i)^\\\\s*PASS\\\\s*$"',
        '</custom_item>',
        '',
    ]

OUTPUT.write_text('\n'.join(parts), encoding='utf-8')
print(f"{OUTPUT.name}: {len(rows)} custom_items gerados")
