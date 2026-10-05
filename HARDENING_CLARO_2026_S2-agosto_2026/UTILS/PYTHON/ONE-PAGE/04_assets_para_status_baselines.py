#!/usr/bin/env python3
"""Atualiza o status de baselines a partir de um snapshot de assets do Tenable."""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import sys
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


TEAM_TAGS = {
    "TI": "CLARO-TI",
    "BCC": "BCC",
    "StarOne": "STAR-ONE",
    "Engenharia Video": "ENG-VIDEO",
}
REQUIRED_HEADERS = {
    "display_operating_system",
    "operating_systems",
    "has_agent",
    "tags",
}
NUMERIC_COLUMNS = {
    "ativos",
    "em_uso",
    "producao",
    "sox",
    "com_agent",
    "escaneados",
    "nao_escaneados",
}


@dataclass
class SnapshotProfile:
    chave: str
    dominio: str
    familia: str
    versao: str
    perfil: str
    ativos: int = 0
    em_uso: int = 0
    com_agent: int = 0
    escaneados: int = 0
    nao_escaneados: int = 0


def load_classifier(script: Path):
    """Carrega a classificação já mantida no 01_EXPORT_TO_DASH.py."""
    spec = importlib.util.spec_from_file_location("export_to_dash", script)
    if spec is None or spec.loader is None:
        raise ValueError(f"Não foi possível carregar o classificador: {script}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def truthy(value: str) -> bool:
    return str(value or "").strip().casefold() in {"true", "sim", "yes", "1"}


def parse_tags(raw: str) -> list[dict[str, object]]:
    try:
        value = json.loads(raw or "[]")
    except json.JSONDecodeError:
        return []
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def team_from_tags(tags: list[dict[str, object]]) -> str | None:
    teams = {
        TEAM_TAGS[str(tag.get("category", "")).strip()]
        for tag in tags
        if str(tag.get("category", "")).strip() in TEAM_TAGS
    }
    if len(teams) > 1:
        raise ValueError(f"Asset associado a mais de uma equipe: {', '.join(sorted(teams))}")
    return next(iter(teams), None)


def was_scanned(row: dict[str, str]) -> bool:
    return bool(
        row.get("last_authenticated_scan_time", "").strip()
        or row.get("last_licensed_scan_time", "").strip()
    )


def read_snapshot(path: Path, classifier) -> tuple[dict[str, dict[str, SnapshotProfile]], Counter]:
    grouped: dict[str, dict[str, SnapshotProfile]] = defaultdict(dict)
    audit = Counter()
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = set(reader.fieldnames or [])
        missing = REQUIRED_HEADERS - headers
        if missing:
            raise ValueError(f"CSV sem colunas obrigatórias: {', '.join(sorted(missing))}")

        for row in reader:
            audit["assets_csv"] += 1
            team = team_from_tags(parse_tags(row.get("tags", "")))
            if team is None:
                audit["sem_equipe"] += 1
                continue
            audit["com_equipe"] += 1
            audit[f"equipe:{team}"] += 1

            classification = classifier.classify(
                row.get("display_operating_system", ""),
                row.get("operating_systems", ""),
            )
            key = "|".join(
                (
                    classification.dominio,
                    classification.familia,
                    classification.versao_principal,
                )
            )
            profile = grouped[team].get(key)
            if profile is None:
                version = classification.versao_principal
                profile_name = (
                    f"{classification.familia} {version}"
                    if version != "Não identificada"
                    else f"{classification.familia} / versão não identificada"
                )
                profile = SnapshotProfile(
                    chave=key,
                    dominio=classification.dominio,
                    familia=classification.familia,
                    versao=version,
                    perfil=profile_name,
                )
                grouped[team][key] = profile

            profile.ativos += 1
            profile.em_uso += 1
            if truthy(row.get("has_agent", "")):
                profile.com_agent += 1
            if was_scanned(row):
                profile.escaneados += 1
            else:
                profile.nao_escaneados += 1
    return grouped, audit


def headers_of(sheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }


def copy_row_style(sheet, source_row: int, target_row: int) -> None:
    for source in sheet[source_row]:
        target = sheet.cell(target_row, source.column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def update_sheet(sheet, profiles: dict[str, SnapshotProfile], exported_at: str) -> Counter:
    headers = headers_of(sheet)
    required = {"chave", "dominio", "familia", "versao", "perfil", "ativos", "em_uso"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Aba {sheet.title} sem colunas: {', '.join(sorted(missing))}")

    existing: dict[str, int] = {}
    for row_number in range(2, sheet.max_row + 1):
        key = str(sheet.cell(row_number, headers["chave"]).value or "").strip()
        if key:
            existing[key] = row_number

    result = Counter()
    for key, row_number in existing.items():
        profile = profiles.get(key)
        # O inventário (ativos/em_uso/produção/SOX) vem da CMDB/Jira e permanece
        # intacto. O snapshot do Tenable atualiza somente agent e scans.
        for field in ("com_agent", "escaneados", "nao_escaneados"):
            if field in headers:
                sheet.cell(row_number, headers[field]).value = getattr(profile, field, 0) if profile else 0
        if "exportado_em" in headers:
            sheet.cell(row_number, headers["exportado_em"]).value = exported_at
        result["atualizados" if profile else "sem_snapshot"] += 1

    for key in set(profiles) - set(existing):
        profile = profiles[key]
        result["perfis_fora_modelo"] += 1
        result["assets_fora_modelo"] += profile.ativos
        result["agents_fora_modelo"] += profile.com_agent

    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    return result


def write_summary(workbook, audit: Counter, changes: dict[str, Counter], source: Path, template: Path) -> None:
    if "RESUMO" in workbook.sheetnames:
        del workbook["RESUMO"]
    sheet = workbook.create_sheet("RESUMO", 0)
    rows = [
        ("Atualização do status de baselines", ""),
        ("Snapshot", source.name),
        ("Modelo anterior", template.name),
        ("Gerado em", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Assets no CSV", audit["assets_csv"]),
        ("Assets distribuídos por equipe", audit["com_equipe"]),
        ("Assets sem tag de equipe (não distribuídos)", audit["sem_equipe"]),
        ("", ""),
        ("Observação", "Ativos, em uso, produção e SOX vêm da CMDB/Jira e foram preservados do arquivo anterior."),
        ("Atualização", "O snapshot do Tenable atualizou somente com_agent, escaneados e nao_escaneados."),
        ("Regra", "Cada asset entra na aba indicada pela categoria TI, BCC, StarOne ou Engenharia Video em tags."),
        ("", ""),
        ("Equipe", "Assets snapshot | perfis atualizados | perfis sem snapshot | perfis fora do modelo"),
    ]
    for team in TEAM_TAGS.values():
        item = changes.get(team, Counter())
        rows.append(
            (
                team,
                f"{audit[f'equipe:{team}']} | {item['atualizados']} | {item['sem_snapshot']} | {item['perfis_fora_modelo']}",
            )
        )
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 100
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Gera o XLSX de status a partir do snapshot de assets.")
    parser.add_argument("snapshot", type=Path, help="CSV de assets")
    parser.add_argument("anterior", type=Path, help="XLSX de status anterior usado como modelo")
    parser.add_argument("-o", "--saida", type=Path, required=True, help="Novo XLSX de saída")
    parser.add_argument(
        "--classificador",
        type=Path,
        default=Path(__file__).with_name("01_EXPORT_TO_DASH.py"),
        help="Script que contém a classificação de sistemas operacionais",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    for path in (args.snapshot, args.anterior, args.classificador):
        if not path.is_file():
            raise SystemExit(f"Arquivo não encontrado: {path}")
    try:
        classifier = load_classifier(args.classificador)
        profiles, audit = read_snapshot(args.snapshot, classifier)
        workbook = load_workbook(args.anterior)
        exported_at = datetime.now().astimezone().isoformat(timespec="seconds")
        changes = {}
        for team in TEAM_TAGS.values():
            if team not in workbook.sheetnames:
                raise ValueError(f"Aba ausente no modelo: {team}")
            changes[team] = update_sheet(workbook[team], profiles.get(team, {}), exported_at)
        write_summary(workbook, audit, changes, args.snapshot, args.anterior)
        workbook.save(args.saida)
    except (ValueError, OSError, json.JSONDecodeError) as error:
        raise SystemExit(f"Erro: {error}") from error

    print(f"{args.saida}: {audit['com_equipe']} assets distribuídos; {audit['sem_equipe']} sem equipe")
    for team in TEAM_TAGS.values():
        item = changes[team]
        print(
            f"  {team}: {audit[f'equipe:{team}']} assets no snapshot; "
            f"{item['atualizados']} perfis casados; "
            f"{item['agents_fora_modelo']} agents fora do modelo"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
