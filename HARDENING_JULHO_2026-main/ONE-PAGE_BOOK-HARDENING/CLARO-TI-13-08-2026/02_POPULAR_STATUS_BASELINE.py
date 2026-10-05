#!/usr/bin/env python3
"""Popula o status_baselines_tenable.xlsx com o export de assets do Tenable."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


EQUIPES = {
    "TI": "CLARO-TI", "BCC": "BCC", "StarOne": "STAR-ONE",
    "Engenharia Video": "ENG-VIDEO",
}
METRICAS = (
    "ativos", "em_uso", "com_agent", "sem_agent", "escaneados",
    "nao_escaneados",
)


def verdadeiro(valor: object) -> bool:
    return str(valor or "").strip().casefold() in {"true", "1", "sim", "yes"}


def equipes(tags_json: str) -> set[str]:
    try:
        tags = json.loads(tags_json or "[]")
    except (json.JSONDecodeError, TypeError):
        return set()
    categorias = {str(tag.get("category") or "").strip() for tag in tags}
    return {destino for origem, destino in EQUIPES.items() if origem in categorias}


def versao(texto: str, padrao: str) -> str | None:
    resultado = re.search(padrao, texto, re.I)
    return resultado.group("v") if resultado else None


def perfil(texto: str) -> str | None:
    texto = " ".join((texto or "").split())

    # Família RHEL: todas as distribuições compatíveis alimentam RHEL N.
    familia_rhel = re.search(
        r"(?:red\s+hat\s+enterprise\s+linux|\brhel\b|centos|oracle\s+linux|"
        r"almalinux|alma\s+linux|rocky\s+linux)", texto, re.I,
    )
    if familia_rhel:
        v = versao(
            texto,
            r"(?:red\s+hat\s+enterprise\s+linux|rhel|centos(?:\s+linux)?|"
            r"oracle\s+linux(?:\s+server)?|almalinux|alma\s+linux|rocky\s+linux)"
            r"(?:\s+server|\s+release)?\s+(?P<v>[6-9])(?:\.\d+)?",
        )
        if not v:
            kernel = versao(texto, r"linux\s+kernel\s+(?P<v>[345])(?:\.x|\.\d+)")
            v = {"3": "7", "4": "8", "5": "9"}.get(kernel or "")
        if v in {"6", "7", "8", "9"}:
            return f"Linux|RHEL|{v}"

    v = versao(texto, r"debian(?:\s+(?:linux|gnu/linux))?(?:\s+release)?\s+(?P<v>\d+)")
    if v and int(v) >= 11:
        return f"Linux|Debian|{v}"

    v = versao(texto, r"ubuntu(?:\s+linux)?(?:\s+release)?\s+(?P<v>\d+)")
    if v and int(v) >= 16:
        return f"Linux|Ubuntu|{v}"

    v = versao(
        texto,
        r"windows\s+server\s+(?P<v>2008|2012|2016|2019|2022|2025)",
    )
    if v and int(v) > 2008:
        return f"Windows|Windows Server|{v}"

    if re.search(r"windows", texto, re.I) and not re.search(r"server", texto, re.I):
        v = versao(texto, r"windows(?:,)?\s+(?P<v>10|11)\b")
        if v:
            return f"Windows|Windows Desktop|{v}"
    return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("assets", type=Path)
    parser.add_argument("status", type=Path, nargs="?", default=Path("status_baselines_tenable.xlsx"))
    args = parser.parse_args()

    workbook = load_workbook(args.status)
    alvos = {}
    for worksheet in workbook.worksheets:
        colunas = {str(c.value): c.column for c in worksheet[1] if c.value}
        for linha in range(2, worksheet.max_row + 1):
            chave = str(worksheet.cell(linha, colunas["chave"]).value or "")
            alvos[worksheet.title, chave] = (worksheet, linha, colunas)
            for campo in METRICAS:
                worksheet.cell(linha, colunas[campo], 0)

    numeros = Counter()
    resumo = Counter()
    with args.assets.open(encoding="utf-8-sig", newline="") as arquivo:
        for ativo in csv.DictReader(arquivo):
            resumo["lidos"] += 1
            chave = perfil(ativo.get("display_operating_system", ""))
            times = equipes(ativo.get("tags", ""))
            if not chave:
                resumo["fora_catalogo"] += 1
                continue
            if not times:
                resumo["sem_equipe"] += 1
                continue
            agent = verdadeiro(ativo.get("has_agent"))
            scan = agent and bool((ativo.get("last_authenticated_scan_time") or "").strip())
            for time in times:
                if (time, chave) not in alvos:
                    resumo["sem_linha"] += 1
                    continue
                numeros[time, chave, "ativos"] += 1
                numeros[time, chave, "em_uso"] += 1
                numeros[time, chave, "com_agent"] += int(agent)
                numeros[time, chave, "sem_agent"] += int(not agent)
                numeros[time, chave, "escaneados"] += int(scan)
                numeros[time, chave, "nao_escaneados"] += int(agent and not scan)
                resumo["preenchidos"] += 1

    for (time, chave), (worksheet, linha, colunas) in alvos.items():
        for campo in METRICAS:
            worksheet.cell(linha, colunas[campo], numeros[time, chave, campo])
    workbook.save(args.status)
    print(f"Atualizado: {args.status}")
    print(
        f"{resumo['lidos']} lidos; {resumo['preenchidos']} preenchidos; "
        f"{resumo['fora_catalogo']} fora do catálogo; {resumo['sem_equipe']} sem equipe"
    )


if __name__ == "__main__":
    main()
