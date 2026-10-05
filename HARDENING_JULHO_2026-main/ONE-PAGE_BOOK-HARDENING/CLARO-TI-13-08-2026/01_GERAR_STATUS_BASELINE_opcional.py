#!/usr/bin/env python3
"""Gera o catálogo canônico de baselines do pipeline Tenable Pure."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


EQUIPES = ("CLARO-TI", "BCC", "STAR-ONE", "ENG-VIDEO")
CAMPOS = (
    "chave", "dominio", "familia", "versao", "perfil", "ativos",
    "em_uso", "producao", "sox", "com_agent", "sem_agent", "escaneados",
    "nao_escaneados", "status", "baseline", "responsavel", "observacao",
)

CATALOGO = {
    "RHEL": ("Linux", ("6", "7", "8", "9")),
    "Debian": ("Linux", ("11", "12", "13")),
    "Ubuntu": ("Linux", ("16", "18", "20", "22", "24")),
    "Windows Server": ("Windows", ("2012", "2016", "2019", "2022", "2025")),
    "Windows Desktop": ("Windows", ("10", "11")),
}


def main() -> None:
    saida = Path(__file__).with_name("status_baselines_tenable.xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for equipe in EQUIPES:
        worksheet = workbook.create_sheet(equipe)
        worksheet.append(CAMPOS)
        for celula in worksheet[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="185F91")
        for familia, (dominio, versoes) in CATALOGO.items():
            for versao in versoes:
                worksheet.append((
                    f"{dominio}|{familia}|{versao}", dominio, familia, versao,
                    f"{familia} {versao}", 0, 0, 0, 0, 0, 0, 0, 0,
                    "Criado", "", "", "",
                ))
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for indice, largura in enumerate(
            (32, 14, 22, 12, 28, 11, 11, 11, 9, 13, 13, 13, 18, 15, 24, 28, 38), 1
        ):
            worksheet.column_dimensions[get_column_letter(indice)].width = largura
    workbook.save(saida)
    print(f"Gerado: {saida} — {sum(len(v[1]) for v in CATALOGO.values())} perfis por equipe")


if __name__ == "__main__":
    main()
