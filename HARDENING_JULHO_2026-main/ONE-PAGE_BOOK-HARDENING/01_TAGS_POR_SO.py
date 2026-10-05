#!/usr/bin/env python3
"""Consolida ativos por equipe e sistema operacional."""

import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


EQUIPES = {
    "BCC": "BCC",
    "Engenharia Video": "ENG Video",
    "TI": "CLARO-TI",
    "StarOne": "STAR ONE",
}

ABAS = {
    "BCC": "BCC",
    "ENG Video": "ENG-VIDEO",
    "CLARO-TI": "CLARO-TI",
    "STAR ONE": "STAR-ONE",
}


def identificar_perfil(sistema: str) -> str | None:
    regras = (
        (r"Amazon Linux.*\b2(?:\D|$)", "Amazon Linux 2"),
        (r"CentOS(?: Linux)?(?: release)?\D*([4578])", "CentOS {}"),
        (r"Oracle Linux.*?\b([4-9])(?:\D|$)", "Oracle Linux {}"),
        (r"Red Hat Enterprise Linux.*?\b([4-9])(?:\D|$)", "RHEL {}"),
        (r"Ubuntu(?: Linux)?\D*(16|20|22)(?:\D|$)", "Ubuntu {}"),
        (r"SUSE.*?\b(10|11|12)(?:\D|$)", "SUSE {}"),
        (r"AIX\D*6(?:\D|$)", "AIX 6"),
        (r"HP-UX\D*11(?:\D|$)", "HP-UX 11"),
        (r"Solaris\D*5(?:\D|$)", "Solaris 5"),
        (r"Windows Server\D*(2003|2008 R2|2008|2012 R2|2012|2016|2019|2022|2025)", "Windows Server {}"),
        (r"Windows 10", "Windows Desktop 10"),
        (r"Windows 2000", "Windows Desktop 2000"),
        (r"z/OS\D*2(?:\D|$)", "z/OS 2"),
        (r"VMware ESXi\D*7(?:\D|$)", "VMware ESXi 7"),
    )
    for expressao, perfil in regras:
        resultado = re.search(expressao, sistema, re.IGNORECASE)
        if resultado:
            return perfil.format(*resultado.groups())
    return None


def preencher_planilha(planilha: Path, totais: Counter) -> tuple[int, int]:
    workbook = load_workbook(planilha)
    consolidados = Counter()
    sem_perfil = 0

    perfis = {}
    for equipe, aba in ABAS.items():
        worksheet = workbook[aba]
        cabecalho = {celula.value: celula.column for celula in worksheet[1]}
        perfis[equipe] = {
            worksheet.cell(linha, cabecalho["perfil"]).value
            for linha in range(2, worksheet.max_row + 1)
        }

    for (equipe, sistema), quantidade in totais.items():
        perfil = identificar_perfil(sistema)
        if perfil in perfis[equipe]:
            consolidados[equipe, perfil] += quantidade
        else:
            sem_perfil += quantidade

    preenchidos = 0
    for equipe, aba in ABAS.items():
        worksheet = workbook[aba]
        cabecalho = {celula.value: celula.column for celula in worksheet[1]}
        for linha in range(2, worksheet.max_row + 1):
            perfil = worksheet.cell(linha, cabecalho["perfil"]).value
            quantidade = consolidados[equipe, perfil]
            worksheet.cell(linha, cabecalho["ativos"], quantidade)
            worksheet.cell(linha, cabecalho["em_uso"], quantidade)
            preenchidos += quantidade

    workbook.save(planilha)
    return preenchidos, sem_perfil


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("entrada", type=Path)
    parser.add_argument("-o", "--saida", type=Path, default=Path("equipes_por_so.csv"))
    parser.add_argument(
        "-p",
        "--planilha",
        type=Path,
        default=Path("status_baselines_todos_2026-08-07_preenchido.xlsx"),
    )
    args = parser.parse_args()

    totais = Counter()
    with args.entrada.open(encoding="utf-8-sig", newline="") as arquivo:
        for linha in csv.DictReader(arquivo):
            tags = json.loads(linha.get("tags") or "[]")
            categorias = {tag.get("category") for tag in tags}
            sistema = (linha.get("display_operating_system") or "Não identificado").strip()

            for categoria, equipe in EQUIPES.items():
                if categoria in categorias:
                    totais[equipe, sistema] += 1

    with args.saida.open("w", encoding="utf-8-sig", newline="") as arquivo:
        campos = ("equipe", "sistema_operacional", "quantidade")
        escritor = csv.DictWriter(arquivo, fieldnames=campos, delimiter=";")
        escritor.writeheader()
        for (equipe, sistema), quantidade in sorted(totais.items()):
            escritor.writerow({
                "equipe": equipe,
                "sistema_operacional": sistema,
                "quantidade": quantidade,
            })

    preenchidos, sem_perfil = preencher_planilha(args.planilha, totais)
    print(f"Gerado: {args.saida} ({len(totais)} combinações)")
    print(f"Planilha preenchida: {args.planilha} ({preenchidos} ativos)")
    print(f"Sem perfil correspondente: {sem_perfil}")


if __name__ == "__main__":
    main()
